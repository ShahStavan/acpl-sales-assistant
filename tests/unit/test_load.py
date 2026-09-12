"""Tests for ``load.py``: date formats, dtypes, document tagging, playbook bool."""

from __future__ import annotations

import calendar
from pathlib import Path
import tempfile

from docx import Document as DocxDocument
import openpyxl
import pandas as pd
import pytest

from acpl_assistant.prepare.load import (
    RawPack,
    _month_vocab,
    _resolve_month_mentions,
    _tag_doc,
    load_raw_pack,
    read_dim_distributor,
    read_dim_geo,
    read_dim_sku,
    read_documents,
    read_fact_primary_sales,
    read_fact_targets,
    read_playbook,
    read_promotions,
    read_stockouts,
)

# ---------------------------------------------------------------------------
# Fixtures — small synthetic data pack on-disk
# ---------------------------------------------------------------------------


@pytest.fixture()
def tmp_data(tmp_path: Path) -> Path:
    d = tmp_path / "data"
    d.mkdir()
    (d / "documents").mkdir()

    # fact_primary_sales — ISO week_start
    pd.DataFrame(
        {
            "week_start": ["2025-07-01", "2025-07-08"],
            "sku_code": ["BS-0101", "BS-0101"],
            "territory_code": ["T-N1", "T-N1"],
            "units": [100, 200],
            "value_inr": [1234.56, 2469.12],
        }
    ).to_csv(d / "fact_primary_sales.csv", index=False)

    # fact_targets — YYYY-MM month
    pd.DataFrame(
        {
            "month": ["2025-07", "2025-08"],
            "brand_name": ["GlucoJoy", "GlucoJoy"],
            "region_name": ["North", "North"],
            "target_value_inr": [1_000_000, 1_100_000],
        }
    ).to_csv(d / "fact_targets.csv", index=False)

    # stockouts — ISO week_start, item_code (not sku_code)
    pd.DataFrame(
        {
            "distributor_id": ["D001", "D002"],
            "item_code": ["BS-0101", "BS-0201"],
            "week_start": ["2025-07-01", "2025-07-08"],
            "region": ["North", "South"],
            "days_out_of_stock": [3, 5],
        }
    ).to_csv(d / "stockouts.csv", index=False)

    # promotions — DD/MM/YYYY dates, sku column
    pd.DataFrame(
        {
            "promo_id": ["PR-0001", "PR-0002"],
            "sku": ["BS-0101", "BS-0201"],
            "region": ["North", "South"],
            "start_date": ["01/07/2025", "15/07/2025"],
            "end_date": ["28/07/2025", "04/08/2025"],
            "discount_pct": [10, 15],
            "mechanic": ["Price-off", "Buy 2 Get 1"],
        }
    ).to_csv(d / "promotions.csv", index=False)

    # dim_sku
    pd.DataFrame(
        {
            "sku_code": ["BS-0101", "BS-0201", "BV-0101", "SN-0101"],
            "sku_name": ["GlucoJoy 50g", "CremeDelight 50g", "Aqualite 1L", "Namkeen Nation 100g"],
            "brand": ["GlucoJoy", "CremeDelight", "Aqualite", "Namkeen Nation"],
            "category": ["Biscuits", "Biscuits", "Beverages", "Snacks"],
            "pack_size": ["50g", "50g", "1L", "100g"],
            "mrp_inr": [34, 14, 60, 20],
        }
    ).to_csv(d / "dim_sku.csv", index=False)

    # dim_geo
    pd.DataFrame(
        {
            "territory_code": ["T-N1", "T-S1", "T-E1", "T-W1"],
            "territory_name": ["Delhi NCR", "Bengaluru", "Kolkata", "Mumbai"],
            "region": ["North", "South", "East", "West"],
            "state": ["Delhi", "Karnataka", "West Bengal", "Maharashtra"],
        }
    ).to_csv(d / "dim_geo.csv", index=False)

    # dim_distributor
    pd.DataFrame(
        {
            "distributor_id": ["D001", "D002", "D032", "D033"],
            "distributor_name": ["Delhi Sales", "Chennai Corp", "Mumbai A", "Mumbai B"],
            "territory_code": ["T-N1", "T-S1", "T-W1", "T-W1"],
            "city": ["Delhi NCR", "Chennai", "Mumbai", "Mumbai"],
        }
    ).to_csv(d / "dim_distributor.csv", index=False)

    # documents — two tiny .docx
    _write_docx(
        d / "documents" / "visit_note_north.docx",
        "North region visit. CremeDelight underperformed in February.\nD002 reported stock-outs.",
    )
    _write_docx(
        d / "documents" / "distributor_note.docx",
        "West: D032 and D033 have Aqualite issues in May.",
    )

    # action_playbook.xlsx
    _write_playbook(d / "action_playbook.xlsx")

    return d


def _write_docx(path: Path, text: str) -> None:
    doc = DocxDocument()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    doc.save(str(path))


def _write_playbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "playbook"
    ws.append(["rule_id", "condition", "recommendation", "action", "needs_approval"])
    ws.append(["R-01", "cond1", "rec1", "act1", "Yes"])
    ws.append(["R-02", "cond2", "rec2", "act2", "No"])
    wb.save(str(path))


# ---------------------------------------------------------------------------
# Date formats
# ---------------------------------------------------------------------------


class TestDateFormats:
    def test_sales_iso(self, tmp_data: Path) -> None:
        df = read_fact_primary_sales(tmp_data)
        assert pd.api.types.is_datetime64_any_dtype(df["week_start"])
        assert df["week_start"].iloc[0] == pd.Timestamp("2025-07-01")

    def test_targets_ym(self, tmp_data: Path) -> None:
        df = read_fact_targets(tmp_data)
        assert pd.api.types.is_datetime64_any_dtype(df["month"])
        assert df["month"].iloc[0] == pd.Timestamp("2025-07-01")

    def test_stockouts_iso(self, tmp_data: Path) -> None:
        df = read_stockouts(tmp_data)
        assert pd.api.types.is_datetime64_any_dtype(df["week_start"])
        assert df["week_start"].iloc[0] == pd.Timestamp("2025-07-01")

    def test_promotions_dmy(self, tmp_data: Path) -> None:
        df = read_promotions(tmp_data)
        assert pd.api.types.is_datetime64_any_dtype(df["start_date"])
        assert pd.api.types.is_datetime64_any_dtype(df["end_date"])
        # 01/07/2025 must be 1 July, never 7 January
        assert df["start_date"].iloc[0] == pd.Timestamp("2025-07-01")
        assert df["end_date"].iloc[0] == pd.Timestamp("2025-07-28")
        # second row has a month crossing
        assert df["end_date"].iloc[1] == pd.Timestamp("2025-08-04")

    def test_promotions_dmy_january_july(self) -> None:
        """01/07/2025 must become 1 July 2025, never 7 January."""
        with tempfile.NamedTemporaryFile(mode="w", suffix=".csv", delete=False, newline="") as f:
            f.write("promo_id,sku,region,start_date,end_date,discount_pct,mechanic\n")
            f.write("PR-T,BS-0101,North,01/07/2025,01/07/2025,10,Price-off\n")
            p = f.name
        try:
            df = read_promotions(Path(p).parent)
        except FileNotFoundError:
            # Read via raw pandas for standalone CSV test
            df = pd.read_csv(
                p,
                dtype={
                    "promo_id": "str",
                    "sku": "str",
                    "region": "str",
                    "discount_pct": "int64",
                    "mechanic": "str",
                },
                parse_dates=["start_date", "end_date"],
                date_format="%d/%m/%Y",
            )
        assert df["start_date"].iloc[0] == pd.Timestamp("2025-07-01"), (
            f"Got {df['start_date'].iloc[0]} instead of 2025-07-01"
        )
        Path(p).unlink(missing_ok=True)


# ---------------------------------------------------------------------------
# Dtypes
# ---------------------------------------------------------------------------


class TestDtypes:
    def test_sales_dtypes(self, tmp_data: Path) -> None:
        df = read_fact_primary_sales(tmp_data)
        assert pd.api.types.is_datetime64_any_dtype(df["week_start"])
        assert pd.api.types.is_string_dtype(df["sku_code"])
        assert pd.api.types.is_string_dtype(df["territory_code"])
        assert pd.api.types.is_integer_dtype(df["units"])
        assert pd.api.types.is_float_dtype(df["value_inr"])

    def test_targets_dtypes(self, tmp_data: Path) -> None:
        df = read_fact_targets(tmp_data)
        assert pd.api.types.is_datetime64_any_dtype(df["month"])
        assert pd.api.types.is_string_dtype(df["brand_name"])
        assert pd.api.types.is_string_dtype(df["region_name"])
        assert pd.api.types.is_integer_dtype(df["target_value_inr"])

    def test_stockouts_dtypes(self, tmp_data: Path) -> None:
        df = read_stockouts(tmp_data)
        assert pd.api.types.is_string_dtype(df["distributor_id"])
        assert pd.api.types.is_string_dtype(df["item_code"])
        assert pd.api.types.is_datetime64_any_dtype(df["week_start"])
        assert pd.api.types.is_string_dtype(df["region"])
        assert pd.api.types.is_integer_dtype(df["days_out_of_stock"])

    def test_promotions_dtypes(self, tmp_data: Path) -> None:
        df = read_promotions(tmp_data)
        assert pd.api.types.is_string_dtype(df["promo_id"])
        assert pd.api.types.is_string_dtype(df["sku"])
        assert pd.api.types.is_string_dtype(df["region"])
        assert pd.api.types.is_datetime64_any_dtype(df["start_date"])
        assert pd.api.types.is_datetime64_any_dtype(df["end_date"])
        assert pd.api.types.is_integer_dtype(df["discount_pct"])
        assert pd.api.types.is_string_dtype(df["mechanic"])

    def test_dim_sku_dtypes(self, tmp_data: Path) -> None:
        df = read_dim_sku(tmp_data)
        assert pd.api.types.is_string_dtype(df["sku_code"])
        assert pd.api.types.is_string_dtype(df["sku_name"])
        assert pd.api.types.is_string_dtype(df["brand"])
        assert pd.api.types.is_string_dtype(df["category"])
        assert pd.api.types.is_string_dtype(df["pack_size"])
        assert pd.api.types.is_integer_dtype(df["mrp_inr"])

    def test_dim_geo_dtypes(self, tmp_data: Path) -> None:
        df = read_dim_geo(tmp_data)
        assert pd.api.types.is_string_dtype(df["territory_code"])
        assert pd.api.types.is_string_dtype(df["territory_name"])
        assert pd.api.types.is_string_dtype(df["region"])
        assert pd.api.types.is_string_dtype(df["state"])

    def test_dim_distributor_dtypes(self, tmp_data: Path) -> None:
        df = read_dim_distributor(tmp_data)
        assert pd.api.types.is_string_dtype(df["distributor_id"])
        assert pd.api.types.is_string_dtype(df["distributor_name"])
        assert pd.api.types.is_string_dtype(df["territory_code"])
        assert pd.api.types.is_string_dtype(df["city"])


# ---------------------------------------------------------------------------
# Document tagging
# ---------------------------------------------------------------------------


class TestDocumentTagging:
    def test_tag_doc_finds_brands(self) -> None:
        result = _tag_doc(
            "CremeDelight is great. GlucoJoy is also good.",
            {"CremeDelight", "GlucoJoy", "Aqualite"},
            set(),
            set(),
        )
        assert result["brands"] == ["CremeDelight", "GlucoJoy"]

    def test_tag_doc_finds_regions(self) -> None:
        result = _tag_doc(
            "Sales in North and East were strong.",
            set(),
            {"North", "South", "East", "West"},
            set(),
        )
        assert result["regions"] == ["East", "North"]

    def test_tag_doc_finds_distributors(self) -> None:
        result = _tag_doc(
            "D032 reported issues. D033 is fine.",
            set(),
            set(),
            {"D032", "D033", "D001"},
        )
        assert result["distributors"] == ["D032", "D033"]

    def test_tag_doc_finds_months(self) -> None:
        result = _tag_doc(
            "Performance in February and April.",
            set(),
            set(),
            set(),
        )
        assert "February" in result["months"]
        assert "April" in result["months"]

    def test_tag_doc_no_match(self) -> None:
        result = _tag_doc("Nothing relevant here.", set(), set(), set())
        assert result["brands"] == []
        assert result["regions"] == []
        assert result["distributors"] == []
        assert result["months"] == []

    def test_documents_integration(self, tmp_data: Path) -> None:
        sku = read_dim_sku(tmp_data)
        geo = read_dim_geo(tmp_data)
        dist = read_dim_distributor(tmp_data)
        docs = read_documents(tmp_data, sku, geo, dist)

        assert len(docs) == 2
        assert "months_resolved" in docs.columns

        note = docs.loc[docs["source_file"] == "visit_note_north.docx"].iloc[0]
        assert "North" in note["regions"]
        assert "CremeDelight" in note["brands"]
        assert "February" in note["months"]
        assert note["months_resolved"] == ""
        assert "D002" in note["distributors"]

        dnote = docs.loc[docs["source_file"] == "distributor_note.docx"].iloc[0]
        assert "West" in dnote["regions"]
        assert "Aqualite" in dnote["brands"]
        assert "May" in dnote["months"]
        assert dnote["months_resolved"] == ""
        assert "D032" in dnote["distributors"]
        assert "D033" in dnote["distributors"]


# ---------------------------------------------------------------------------
# Playbook bool
# ---------------------------------------------------------------------------


class TestPlaybook:
    def test_needs_approval_bool(self, tmp_data: Path) -> None:
        df = read_playbook(tmp_data)
        assert pd.api.types.is_bool_dtype(df["needs_approval"])
        assert df["needs_approval"].iloc[0]
        assert not df["needs_approval"].iloc[1]

    def test_needs_approval_no_mixed_types(self, tmp_data: Path) -> None:
        df = read_playbook(tmp_data)
        assert set(df["needs_approval"].unique()) == {True, False}


# ---------------------------------------------------------------------------
# RawPack / load_raw_pack
# ---------------------------------------------------------------------------


class TestRawPack:
    def test_load_raw_pack_returns_all_sources(self, tmp_data: Path) -> None:
        pack = load_raw_pack(tmp_data)
        assert isinstance(pack, RawPack)
        assert len(pack.fact_primary_sales) == 2
        assert len(pack.fact_targets) == 2
        assert len(pack.stockouts) == 2
        assert len(pack.promotions) == 2
        assert len(pack.dim_sku) == 4
        assert len(pack.dim_geo) == 4
        assert len(pack.dim_distributor) == 4
        assert len(pack.documents) == 2
        assert len(pack.playbook) == 2

    def test_load_raw_pack_default_dir(self) -> None:
        pack = load_raw_pack()
        assert isinstance(pack, RawPack)
        assert len(pack.dim_sku) == 120
        assert len(pack.dim_geo) == 12
        assert len(pack.dim_distributor) == 40


# ---------------------------------------------------------------------------
# Month vocabulary
# ---------------------------------------------------------------------------


class TestMonthVocab:
    def test_contains_all_months(self) -> None:
        v = _month_vocab()
        for i in range(1, 13):
            assert calendar.month_name[i] in v
            assert calendar.month_abbr[i] in v

    def test_has_no_extras(self) -> None:
        v = _month_vocab()
        assert "NotAMonth" not in v


# ---------------------------------------------------------------------------
# Whole-word matching (regression for "Market" → "Mar")
# ---------------------------------------------------------------------------


class TestWholeWordMatching:
    def test_market_does_not_tag_march(self) -> None:
        result = _tag_doc(
            "Market visit went well.",
            set(),
            set(),
            set(),
        )
        assert "Mar" not in result["months"], "Market must not match Mar via substring"


# ---------------------------------------------------------------------------
# Month year resolution
# ---------------------------------------------------------------------------


class TestMonthYearResolution:
    def test_february_resolves_to_feb2026(self) -> None:
        resolved = _resolve_month_mentions(
            "Performance in February 2026 was strong.",
            "some_file.docx",
        )
        assert "2026-02" in resolved

    def test_year_from_filename(self) -> None:
        resolved = _resolve_month_mentions(
            "Sales improved in May.",
            "summary_2026.docx",
        )
        assert "2026-05" in resolved

    def test_no_year_skips(self) -> None:
        resolved = _resolve_month_mentions(
            "Sales improved in May.",
            "summary.docx",
        )
        assert resolved == []

    def test_abbr_and_full_collapse(self) -> None:
        resolved = _resolve_month_mentions(
            "Feb and February both present.",
            "report_2026.docx",
        )
        assert len(resolved) == 1
        assert "2026-02" in resolved
